import unittest
from collections import defaultdict
from pathlib import Path
from statistics import pstdev

from src.engine import PRODUCTS, _number, _code, _file_report, build_plan, load_data

BASE = Path(__file__).resolve().parents[1] / 'data' / 'baseline'  # was /Users/unikamaharjan/Downloads (only worked on one Mac)


def sample_data():
    totals=defaultdict(lambda: defaultdict(float)); days=defaultdict(set); regions=defaultdict(float)
    for rd, qty in [(1,100),(2,120)]:
        totals[rd]['CC-T01']=qty; days[rd]=set(range(1,11))
        regions[(rd,'CC-T01','North')]=qty*.5
        regions[(rd,'CC-T01','South')]=qty*.3
        regions[(rd,'CC-T01','West')]=qty*.2
    return {'rounds':[1,2], 'round_totals':totals, 'round_days':days, 'sales':[],
            'regional_round_sales':regions, 'inventory':{('CC-T01','03'):10,('CC-T01','03N'):5,('CC-T01','03S'):2,('CC-T01','03W'):1},
            'inbound':defaultdict(float,{'CC-T01':3}), 'prices':{'CC-T01':30},
            'price_history':defaultdict(list), 'has_inventory':True,'has_po':True,
            'costs':{'CC-T01':22.95},
            'valuations':[], 'financial':{}, 'max_round':2}


class EngineTests(unittest.TestCase):
    def test_known_example_weighted_forecast_and_mrp_netting(self):
        p=build_plan(sample_data(),3)
        milk=p['rows'][0]
        self.assertEqual(milk['forecast'],112)  # Latest-round-biased 3:5 weights, normalized for two rounds.
        self.assertEqual(milk['md61'],milk['forecast']+milk['buffer'])
        self.assertEqual(milk['mrp'],max(0,milk['md61']-21))

    def test_realistic_provided_round_data(self):
        root=Path(__file__).resolve().parents[1]
        data=load_data(root/'data'/'uploads')
        p=build_plan(data,max(data['complete_rounds'])+1)
        self.assertEqual(data['rounds'],list(range(1,10)))
        self.assertEqual(data['complete_rounds'],list(range(1,9)))
        self.assertEqual(data['incomplete_rounds'],[9])
        self.assertEqual(p['target_round'],9)
        self.assertEqual(p['through_round'],8)
        self.assertAlmostEqual(p['company_value'],1125700.90,places=2)
        freshness={f['name']:f['status'] for f in data['files']}
        self.assertEqual(freshness['Detailed sales'],'Outdated')
        self.assertEqual(freshness['Financial'],'Outdated')
        self.assertEqual([r['product'] for r in p['rows']],[x['name'] for x in PRODUCTS])
        self.assertTrue(all(r['mrp'] is None or r['mrp']>=0 for r in p['rows']))

    def test_invalid_inputs_are_safe(self):
        self.assertEqual(_number('not a number'),0.0)
        self.assertEqual(_number(None,7),7)
        self.assertIsNone(_code('unknown material'))
        self.assertIsNone(_code('CC-T09'))

    def test_report_detection_from_columns(self):
        self.assertEqual(_file_report(BASE/'SalesExportData.xlsx'),'detailed_sales')
        self.assertEqual(_file_report(BASE/'ExportData (1).xlsx'),'inventory')
        self.assertEqual(_file_report(BASE/'purchase order round 8 ExportData.xlsx'),'purchase_orders')

    def test_independent_verification_of_round_9_milk_forecast(self):
        # Independently read raw sales and recompute the 20/30/50 weighted forecast.
        from openpyxl import load_workbook
        src=Path(BASE/'O data for all the rounds.xlsx')
        rows=list(load_workbook(src,data_only=True,read_only=True)['Sales'].values)
        header=rows[0]; qi=header.index('QUANTITY'); pi=header.index('MATERIAL_NUMBER'); ri=header.index('SIM_ROUND')
        byround={r:0 for r in (6,7)}
        for row in rows[1:]:
            if str(row[pi]).endswith('T01') and int(row[ri]) in byround: byround[int(row[ri])]+=float(row[qi])
        # Round 8 detail export is independently aggregated.
        x=load_workbook(BASE/'SalesExportData.xlsx',data_only=True,read_only=True).active
        h, *data=list(x.values); qi=h.index('Quantity'); pi=h.index('Material'); ri=h.index('Round')
        byround[8]=sum(float(row[qi] or 0) for row in data if str(row[pi]).endswith('T01') and int(row[ri])==8)
        expected=round(byround[6]*.2+byround[7]*.3+byround[8]*.5)
        plan=build_plan(load_data(Path(__file__).resolve().parents[1]/'data'/'uploads'),9)
        self.assertEqual(plan['rows'][0]['forecast'],expected)


if __name__=='__main__': unittest.main()
