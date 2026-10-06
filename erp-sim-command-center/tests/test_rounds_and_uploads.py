"""Planning-round behaviour and sales uploads (fixes made after the first two revisions)."""
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from src.engine import build_plan, load_data

SALES = Path(__file__).resolve().parents[1] / "data" / "baseline" / "SalesExportData.xlsx"


def sales_file(folder, keep_round=None, relabel_as=None):
    """Copy of the packaged ZVA05 export, optionally keeping one round and renaming it."""
    wb = openpyxl.load_workbook(SALES); ws = wb.active
    for row in range(ws.max_row, 1, -1):
        if keep_round and str(ws.cell(row, 1).value) != keep_round:
            ws.delete_rows(row)
        elif relabel_as:
            ws.cell(row, 1).value = relabel_as
    out = Path(folder) / "MySalesUpload.xlsx"; wb.save(out); return out


class Rounds(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = load_data()

    def test_past_round_is_a_replay_without_todays_stock(self):
        plan = build_plan(self.data, 4, "auto", 2)
        self.assertTrue(plan["replay"])
        self.assertIsNone(plan["purchase_qty"])      # no stock from after round 8 used for round 4
        self.assertIsNone(plan["policy"])
        self.assertIn("Replay of Round 4", plan["data_warning"])

    def test_kpis_follow_the_selected_round(self):
        r4, r9 = build_plan(self.data, 4, "auto", 2), build_plan(self.data, 9, "auto", 2)
        self.assertEqual((r4["valuation_asof"]["round"], r4["valuation_asof"]["day"]), (3, 10))
        self.assertAlmostEqual(r4["company_value"], 1221777.60, places=1)
        self.assertAlmostEqual(r4["financial"]["profit"], 45816.66, places=1)
        self.assertNotEqual(r4["company_value"], r9["company_value"])
        # Cross-check with the course formula: (profit / rounds) x 8 / (7% + 3%)
        self.assertAlmostEqual(r4["financial"]["profit"] / 3 * 8 / 0.10, r4["company_value"], places=0)

    def test_push_pull_does_not_change_sap_kpis(self):
        a, b = build_plan(self.data, 9, "push", 2), build_plan(self.data, 9, "pull", 2)
        self.assertEqual(a["company_value"], b["company_value"])
        self.assertEqual(a["financial"]["profit"], b["financial"]["profit"])


class SalesUploads(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.base = load_data()

    def tearDown(self):
        shutil.rmtree(self.folder, ignore_errors=True)

    def test_same_file_again_changes_nothing_and_says_so(self):
        shutil.copy(SALES, Path(self.folder) / "SalesExportData.xlsx")
        data = load_data(self.folder)
        self.assertFalse(data["upload_report"][0]["changed"])
        self.assertIn("Nothing new", data["upload_report"][0]["detail"])
        self.assertEqual(build_plan(data, 9)["purchase_qty"], build_plan(self.base, 9)["purchase_qty"])

    def test_partial_upload_does_not_wipe_older_rounds(self):
        sales_file(self.folder, keep_round="01")   # before the fix this replaced the packaged file and round 8 vanished
        data = load_data(self.folder)
        self.assertEqual(data["complete_rounds"], self.base["complete_rounds"])
        self.assertEqual(dict(data["round_totals"][8]), dict(self.base["round_totals"][8]))

    def test_new_round_is_added_and_moves_the_plan_forward(self):
        sales_file(self.folder, keep_round="08", relabel_as="09")   # a full 10-day round 9
        data = load_data(self.folder)
        self.assertTrue(data["upload_report"][0]["changed"])
        self.assertIn(9, data["complete_rounds"])
        plan = build_plan(data, max(data["complete_rounds"]) + 1)
        self.assertEqual(plan["target_round"], 10)
        self.assertEqual(plan["through_round"], 9)


if __name__ == "__main__":
    unittest.main()
